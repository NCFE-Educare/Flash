"""Extract text from PDF, DOCX, PPTX, and TXT files for chat context."""

import base64
import os
from pathlib import Path

# Load .env so MISTRAL_API_KEY is available even when called outside FastAPI startup
_env_path = Path(__file__).parent / ".env"
if _env_path.exists():
    for _line in _env_path.read_text().splitlines():
        _line = _line.strip()
        if _line and not _line.startswith("#") and "=" in _line:
            _k, _, _v = _line.partition("=")
            os.environ[_k.strip()] = _v.strip().strip('"').strip("'")

# If pypdf extracts less than this many chars, treat PDF as image-based and try OCR
MIN_TEXT_CHARS_FOR_PDF = 20


def _extract_pdf_with_ocr(path: Path) -> str:
    """Extract text from image-based (scanned) PDF using Mistral OCR API."""
    api_key = os.environ.get("MISTRAL_API_KEY")
    if not api_key:
        raise RuntimeError("MISTRAL_API_KEY not set — required for image-based PDF OCR")

    from mistralai.client import Mistral

    client = Mistral(api_key=api_key)

    with open(path, "rb") as f:
        b64_pdf = base64.b64encode(f.read()).decode("utf-8")

    ocr_response = client.ocr.process(
        model="mistral-ocr-latest",
        document={
            "type": "document_url",
            "document_url": f"data:application/pdf;base64,{b64_pdf}",
        },
    )

    pages = getattr(ocr_response, "pages", None) or []
    parts = []
    for page in pages:
        md = getattr(page, "markdown", None) or ""
        if md.strip():
            parts.append(md.strip())

    return "\n\n".join(parts)


def extract_text_from_file(file_path: Path) -> str:
    """
    Extract text from PDF, DOCX, PPTX, or TXT.
    For PDFs: tries text extraction first; if empty (image-based), falls back to OCR.
    Returns plain text. Raises ValueError for unsupported formats.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace")

    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                pass

        parts = []
        for page in reader.pages:
            text = page.extract_text()
            if text:
                parts.append(text)
        text = "\n".join(parts) if parts else ""

        if len(text.strip()) < MIN_TEXT_CHARS_FOR_PDF and len(reader.pages) > 0:
            try:
                ocr_text = _extract_pdf_with_ocr(path)
                if ocr_text:
                    return ocr_text
                return "[OCR returned no text — the document may be blank or unsupported]"
            except Exception as e:
                return f"[OCR failed: {e}]"

        return text

    if suffix == ".docx":
        from docx import Document

        doc = Document(path)
        return "\n".join(p.text for p in doc.paragraphs)

    if suffix in (".ppt", ".pptx"):
        from pptx import Presentation

        prs = Presentation(path)
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    parts.append(shape.text)
        return "\n".join(parts)

    raise ValueError(f"Unsupported document format: {suffix}")
